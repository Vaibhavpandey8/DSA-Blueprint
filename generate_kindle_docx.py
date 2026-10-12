import os
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

book_dir = os.path.dirname(os.path.abspath(__file__))
complete_md_path = os.path.join(book_dir, "THE_DSA_BLUEPRINT_COMPLETE_BOOK.md")

with open(complete_md_path, "r", encoding="utf-8") as f:
    lines = f.readlines()

doc = Document()

# Set standard margins
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

in_code_block = False
code_lines = []
in_table = False
table_rows = []

def clean_latex(text):
    text = re.sub(r'\\approx', '≈', text)
    text = re.sub(r'\\le', '≤', text)
    text = re.sub(r'\\ge', '≥', text)
    text = re.sub(r'\\rightarrow', '→', text)
    text = re.sub(r'\\implies', '⟹', text)
    text = re.sub(r'\\times', '×', text)
    text = re.sub(r'\\cdot', '·', text)
    text = re.sub(r'\\oplus', '⊕', text)
    text = re.sub(r'\\log', 'log', text)
    text = re.sub(r'\\sum_\{j=0\}\^\{i\}', 'Σ(j=0..i)', text)
    text = re.sub(r'\\text\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\$([^$]+)\$', r'\1', text)
    return text

def add_formatted_text(p, text):
    text = clean_latex(text)
    # Split by bold and inline code
    # Simple regex parsing for **bold** and `code`
    tokens = re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = p.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('`') and token.endswith('`'):
            run = p.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(3, 105, 161)
        else:
            p.add_run(token)

i = 0
while i < len(lines):
    line = lines[i].rstrip('\r\n')
    
    # Check code block
    if line.startswith('```'):
        if in_code_block:
            # End code block
            in_code_block = False
            code_text = '\n'.join(code_lines)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(code_text)
            run.font.name = 'Consolas'
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(30, 41, 59)
            code_lines = []
        else:
            in_code_block = True
            code_lines = []
        i += 1
        continue
        
    if in_code_block:
        code_lines.append(line)
        i += 1
        continue

    # Check table
    if line.strip().startswith('|') and line.strip().endswith('|'):
        table_rows.append(line.strip())
        i += 1
        continue
    else:
        if table_rows:
            # Process table
            parsed = []
            for tr in table_rows:
                cols = [c.strip() for c in tr.split('|')[1:-1]]
                if cols and not all(set(c).issubset({'-', ':', ' '}) for c in cols):
                    parsed.append(cols)
            if parsed:
                max_cols = max(len(r) for r in parsed)
                t = doc.add_table(rows=len(parsed), cols=max_cols)
                t.style = 'Table Grid'
                for r_idx, r_data in enumerate(parsed):
                    for c_idx, cell_data in enumerate(r_data):
                        cell = t.cell(r_idx, c_idx)
                        cp = cell.paragraphs[0]
                        add_formatted_text(cp, cell_data)
                        if r_idx == 0:
                            for r_run in cp.runs:
                                r_run.bold = True
                            set_cell_background(cell, 'F1F5F9')
            table_rows = []

    # Ignore raw html tags like <div>, </div>
    if line.strip().startswith('<div') or line.strip().startswith('</div'):
        if 'page-break-after: always' in line:
            doc.add_page_break()
        i += 1
        continue

    # Blank line
    if not line.strip():
        i += 1
        continue

    # Headings
    if line.startswith('# '):
        doc.add_heading(clean_latex(line[2:].strip()), level=1)
    elif line.startswith('## '):
        doc.add_heading(clean_latex(line[3:].strip()), level=2)
    elif line.startswith('### '):
        doc.add_heading(clean_latex(line[4:].strip()), level=3)
    elif line.startswith('#### '):
        doc.add_heading(clean_latex(line[5:].strip()), level=4)
    elif line.startswith('> '):
        p = doc.add_paragraph(style='Quote')
        add_formatted_text(p, line[2:].strip())
    elif line.startswith('- ') or line.startswith('* '):
        p = doc.add_paragraph(style='List Bullet')
        add_formatted_text(p, line[2:].strip())
    elif re.match(r'^\d+\.\s+', line):
        p = doc.add_paragraph(style='List Number')
        text_content = re.sub(r'^\d+\.\s+', '', line)
        add_formatted_text(p, text_content)
    elif line.strip() == '---':
        p = doc.add_paragraph()
        p.add_run('—' * 40).font.color.rgb = RGBColor(200, 200, 200)
    else:
        p = doc.add_paragraph()
        add_formatted_text(p, line)

    i += 1

output_docx = os.path.join(book_dir, "THE_DSA_BLUEPRINT_KINDLE.docx")
doc.save(output_docx)
print(f"Generated Kindle Word Docx: {output_docx}")
